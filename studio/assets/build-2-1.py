"""Документи, які потрапляють у кадр уроку 2.1.

Два файли:
  zvernennia-1247.docx  — вхідне звернення мешканки; з нього в кадрі копіюють текст;
  blank-vidpovid.docx   — бланк управління, куди наприкінці вставляють готовий лист.

Чому .docx, а не готова картинка: у кадрі має бути справжній Word, у якому людина
виділяє текст мишею й тисне Ctrl+C. Скріншот документа цього не покаже.

Усі дані вигадані. Числа звірені з текстом уроку
(`prisma/src/content/lessons-module-2.ts`, урок `lyst-vidpovid-na-zvernennya`):
вх. № 1247/03-15 від 12.09.2026, вул. Садова, буд. 12, кв. 5.

Запуск:  python studio/assets/build-2-1.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).parent

# Оформлення за звичаєм українського діловодства: Times New Roman 14,
# поля 30/10/20/20 мм. Саме так виглядає папір, який людина бачить щодня.
FONT = 'Times New Roman'
SIZE = Pt(14)


def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.left_margin, s.right_margin = Cm(3), Cm(1)
    s.top_margin, s.bottom_margin = Cm(2), Cm(2)

    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    return doc


def para(doc, text='', *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, indent=None,
         before=0, after=0, size=None, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent is not None:
        p.paragraph_format.first_line_indent = Cm(indent)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.size = size or SIZE
    return p


def build_zvernennia():
    doc = new_doc()

    # Реєстраційний штамп — те, що ставить канцелярія.
    para(doc, 'Зареєстровано: вх. № 1247/03-15 від 12.09.2026',
         align=WD_ALIGN_PARAGRAPH.LEFT, size=Pt(12), italic=True, after=18)

    # Адресат і заявниця — праворуч, як на справжньому зверненні.
    for line in [
        'Начальнику управління',
        'житлово-комунального господарства',
        'Н-ської міської ради',
        'Петренку І. В.',
    ]:
        para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT)

    para(doc, '', after=10)

    for line in [
        'Коваленко Марії Іванівни,',
        'яка проживає за адресою:',
        'вул. Садова, буд. 12, кв. 5,',
        'м. Н-ськ, 00000',
        'тел. 067 123 45 67',
    ]:
        para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT)

    para(doc, 'Звернення', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, before=28, after=18)

    # Саме ці три абзаци в кадрі виділяють і копіюють у чат. Дослівно з уроку:
    # у зверненні три питання — стан майданчика, небезпека, строки.
    for text in [
        'Прошу вжити заходів щодо дитячого майданчика у дворі нашого будинку. '
        'Гойдалки поламані, огорожа частково відсутня, гумове покриття провалилося.',

        'Майданчиком щодня користуються діти з навколишніх будинків, і через '
        'відсутність огорожі вони вибігають просто на проїжджу частину. Вважаю, '
        'що це створює загрозу для їхнього життя та здоров’я.',

        'Прошу повідомити, які заходи буде вжито та в які строки планується '
        'ремонт майданчика.',
    ]:
        para(doc, text, indent=1.25, after=10)

    para(doc, '', after=20)
    p = para(doc, '12 вересня 2026 р.', align=WD_ALIGN_PARAGRAPH.LEFT)
    p.add_run('\t\t\t\t\tМ. І. Коваленко').font.size = SIZE

    out = HERE / 'zvernennia-1247.docx'
    doc.save(out)
    return out


def build_blank():
    doc = new_doc()

    for line, bold in [
        ('Н-СЬКА МІСЬКА РАДА', True),
        ('УПРАВЛІННЯ ЖИТЛОВО-КОМУНАЛЬНОГО ГОСПОДАРСТВА', True),
    ]:
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=bold)

    para(doc, 'вул. Центральна, 1, м. Н-ськ, 00000, тел. (0000) 00-00-00',
         align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), after=6)

    # Лінія під шапкою бланка.
    para(doc, '_' * 78, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    para(doc, '№ ______________ від ____________', align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    para(doc, 'На № 1247/03-15 від 12.09.2026', align=WD_ALIGN_PARAGRAPH.LEFT, after=18)

    for line in ['Коваленко М. І.', 'вул. Садова, буд. 12, кв. 5,', 'м. Н-ськ, 00000']:
        para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # Порожнє місце: у кадрі сюди вставляють текст із чату.
    para(doc, '', before=24)

    out = HERE / 'blank-vidpovid.docx'
    doc.save(out)
    return out


if __name__ == '__main__':
    for path in (build_zvernennia(), build_blank()):
        print(f'  {path.name}  {path.stat().st_size} байт')
